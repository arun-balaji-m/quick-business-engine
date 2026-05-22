import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.logging import get_logger

logger = get_logger(__name__)


class ExportService:
    def export_csv(self, columns: list[str], rows: list[dict[str, Any]]) -> bytes:
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in columns})
        return output.getvalue().encode("utf-8-sig")  # BOM for Excel compatibility

    def export_excel(
        self,
        columns: list[str],
        rows: list[dict[str, Any]],
        sheet_title: str = "Query Results",
    ) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]  # Excel sheet title max 31 chars

        # Header styling
        header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, col_name in enumerate(columns, start=1):
                value = row.get(col_name)
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-adjust column widths
        for column_cells in ws.columns:
            max_len = 0
            col_letter = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()
